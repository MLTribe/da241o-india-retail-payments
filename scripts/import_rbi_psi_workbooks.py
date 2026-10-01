#!/usr/bin/env python3
"""Extract own-month rows from archived RBI Payment System Indicators workbooks.

RBI serves each monthly release as an XLSX link behind bot protection, so the
workbooks are saved manually from a browser into source_archive/rbi/xlsx/.
Name each file either with its official XLSX basename (as listed in the
release-link index) or with a leading observation month, e.g. 2021-06_PSI.xlsx.
Some RBI links return an HTML table under an .XLSX name; both forms are read.

Run from the repository root: python3 scripts/import_rbi_psi_workbooks.py
Real .xlsx files need openpyxl; HTML-table files need only the standard library.
"""

from __future__ import annotations

import csv
import hashlib
import html
import re
import sys
from datetime import datetime, timezone
from pathlib import Path


REPO = Path(__file__).resolve().parents[1]
YEAR_ROOT = REPO / "by_year"
RBI_ARCHIVE = REPO / "source_archive" / "rbi"
WORKBOOKS = RBI_ARCHIVE / "xlsx"
LINK_INDEX = RBI_ARCHIVE / "rbi_psi_release_links_2022-01_2026-03.csv"
CROSSCHECK = RBI_ARCHIVE / "xlsx_crosscheck.csv"
MANIFEST = REPO / "source_manifest.csv"
RELEASE_ARCHIVE_URL = "https://www.rbi.org.in/Scripts/PSIUserView.aspx"

MONTHS = {name: i for i, name in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

# (section, parent, normalised-label regex) -> output field stem.
# Sections are "txn" before the infrastructure heading and "infra" after it.
# Parent tracks the most recent card/PPI heading so "PoS based"/"Others" resolve.
TRANSACTION_ITEMS: list[tuple[str, str | None, str]] = [
    ("card_total", None, r"^card payments"),
    ("credit_total", None, r"^credit cards?$"),
    ("debit_total", None, r"^debit cards?$"),
    ("credit_pos", "credit_total", r"^(credit cards? - )?pos( based)?$"),
    ("credit_other", "credit_total", r"^(credit cards? - )?others?$"),
    ("debit_pos", "debit_total", r"^(debit cards? - )?pos( based)?$"),
    ("debit_other", "debit_total", r"^(debit cards? - )?others?$"),
    ("ppi_total", None, r"^prepaid payment instruments"),
    ("ppi_wallets", None, r"^wallets$"),
    ("ppi_cards", None, r"^ppi cards$"),
    ("ppi_card_pos", "ppi_cards", r"^(ppi cards? - )?pos( based)?$"),
    ("ppi_card_other", "ppi_cards", r"^(ppi cards? - )?others?$"),
    ("rtgs", None, r"^rtgs$"),
    ("neft", None, r"^neft$"),
    ("aeps_fund_transfer", None, r"^aeps \(fund transfers?\)"),
    ("netc", None, r"^netc( \(linked to bank account\))?$"),
    ("atm_cash_withdrawal", "cash", r"^(cash withdrawal (at|through) )?atms?$"),
]
INFRA_ITEMS: list[tuple[str, str]] = [
    ("credit_cards_outstanding", r"^credit cards?$"),
    ("debit_cards_outstanding", r"^debit cards?$"),
    ("atms", r"^atms?$"),
    ("pos_terminals", r"^pos terminals?$"),
    ("bharat_qr_codes", r"^bharat qr"),
    ("upi_qr_codes", r"^upi qr"),
]
CARD_PPI_FIELDS = [
    "card_total", "credit_total", "credit_pos", "credit_other", "debit_total",
    "debit_pos", "debit_other", "ppi_total", "ppi_wallets", "ppi_cards",
    "ppi_card_pos", "ppi_card_other",
]
OTHER_RAIL_FIELDS = ["rtgs", "neft", "aeps_fund_transfer", "netc", "atm_cash_withdrawal"]


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def normalise_label(text: str) -> str:
    text = clean(text).lower()
    text = re.sub(r"^[\d.]+\s*", "", text)          # "2.1.1 " numbering
    text = re.sub(r"[$#@*^]+", "", text)            # footnote markers
    text = text.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text).strip(" :-")


def number(value: object) -> float | None:
    text = clean(value).replace(",", "")
    if text in ("", "-", "--", "NA", "N.A.", "n.a."):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_month_label(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    text = clean(value).lower().replace("'", "-").replace(".", "")
    match = re.match(r"^([a-z]{3})[a-z]*[- ]?(\d{2}|\d{4})$", text)
    if not match or match.group(1) not in MONTHS:
        return None
    year = int(match.group(2))
    year += 2000 if year < 100 else 0
    return f"{year:04d}-{MONTHS[match.group(1)]:02d}"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_grid(path: Path) -> list[list[object]]:
    head = path.read_bytes()[:4]
    if head.startswith(b"PK"):
        from openpyxl import load_workbook  # only real XLSX files need it
        workbook = load_workbook(path, read_only=True, data_only=True)
        grid = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
        workbook.close()
        return grid
    text = path.read_text(encoding="utf-8", errors="replace")
    if "TSPD" in text and "<table" not in text.lower():
        raise ValueError(f"{path.name} is RBI's bot-check page, not a workbook; "
                         "save the file again from a browser")
    grid = []
    for row in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.S | re.I):
        cells: list[object] = []
        for attrs, body in re.findall(r"<t[dh]([^>]*)>(.*?)</t[dh]>", row, re.S | re.I):
            value = html.unescape(re.sub(r"<[^>]+>", " ", body))
            span = re.search(r'colspan\s*=\s*"?(\d+)', attrs, re.I)
            cells.extend([clean(value)] * (int(span.group(1)) if span else 1))
        grid.append(cells)
    return grid


def month_columns(grid: list[list[object]], month: str) -> tuple[int, int | None]:
    """Return (volume column, value column) for the observation month."""
    for row in grid:
        hits = [i for i, cell in enumerate(row) if parse_month_label(cell) == month]
        if hits:
            return hits[0], hits[1] if len(hits) > 1 else None
    raise ValueError(f"No column header for {month}")


def label_of(row: list[object], first_numeric: int) -> str:
    texts = [clean(cell) for cell in row[:first_numeric] if clean(cell)
             and number(cell) is None]
    return normalise_label(texts[-1]) if texts else ""


def extract(path: Path, month: str) -> tuple[dict[str, float], list[str]]:
    grid = read_grid(path)
    vol_col, val_col = month_columns(grid, month)
    found: dict[str, float] = {}
    section, parent = "txn", None
    for row in grid:
        if not row:
            continue
        whole = normalise_label(" ".join(clean(c) for c in row))
        if "payment system infrastructure" in whole:
            section, parent = "infra", None
            continue
        if section == "txn" and re.search(r"^cash withdrawal", whole):
            parent = "cash"
        label = label_of(row, min(vol_col, len(row)))
        if not label:
            continue
        volume = number(row[vol_col]) if vol_col < len(row) else None
        value = number(row[val_col]) if val_col is not None and val_col < len(row) else None
        if section == "infra":
            for field, pattern in INFRA_ITEMS:
                if field not in found and volume is not None and re.search(pattern, label):
                    found[field] = volume
            continue
        for field, needed_parent, pattern in TRANSACTION_ITEMS:
            if (f"{field}_volume_lakh" in found or not re.search(pattern, label)
                    or (needed_parent is not None and parent != needed_parent)):
                continue
            if volume is not None:
                found[f"{field}_volume_lakh"] = volume
            if value is not None:
                found[f"{field}_value_crore"] = value
            break
        for heading in ("credit_total", "debit_total", "ppi_cards"):
            if re.search(dict((f, p) for f, _, p in TRANSACTION_ITEMS)[heading], label):
                parent = heading
    expected = ([f"{f}_{u}" for f in CARD_PPI_FIELDS + OTHER_RAIL_FIELDS
                 for u in ("volume_lakh", "value_crore")]
                + [field for field, _ in INFRA_ITEMS])
    return found, [field for field in expected if field not in found]


def workbook_months() -> dict[str, tuple[Path, str]]:
    by_name: dict[str, str] = {}
    if LINK_INDEX.exists():
        with LINK_INDEX.open(newline="", encoding="utf-8") as stream:
            for row in csv.DictReader(stream):
                by_name[row["official_xlsx_url"].rsplit("/", 1)[-1].lower()] = row["month"]
    mapping: dict[str, tuple[Path, str]] = {}
    for path in sorted(WORKBOOKS.glob("*")):
        if path.suffix.lower() not in (".xlsx", ".xls", ".html", ".htm"):
            continue
        prefix = re.match(r"^(20\d{2}-\d{2})", path.name)
        month = prefix.group(1) if prefix else by_name.get(path.name.lower())
        if month is None:
            print(f"skip {path.name}: no month prefix and not in the link index",
                  file=sys.stderr)
            continue
        if month in mapping:
            raise ValueError(f"Two workbooks for {month}: {mapping[month][0].name}, {path.name}")
        mapping[month] = (path, sha256(path))
    return mapping


def read_year_files(name: str) -> dict[str, dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    for path in sorted(YEAR_ROOT.glob(f"*/{name}")):
        with path.open(newline="", encoding="utf-8-sig") as stream:
            for row in csv.DictReader(stream):
                rows[row["month"][:7]] = row
    return rows


def write_year_files(name: str, rows: dict[str, dict[str, object]], fields: list[str]) -> None:
    by_year: dict[str, list[dict[str, object]]] = {}
    for month in sorted(rows):
        by_year.setdefault(month[:4], []).append(rows[month])
    for year, items in by_year.items():
        output = YEAR_ROOT / year / name
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(items)


def update_manifest(archived: dict[str, tuple[Path, str]]) -> None:
    with MANIFEST.open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        fields = list(reader.fieldnames or [])
        current = {row["source_id"]: row for row in reader}
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    for month, (path, digest) in archived.items():
        source_id = f"RBI_PSI_XLSX_{month.replace('-', '_')}"
        current[source_id] = {
            "source_id": source_id, "owner": "RBI",
            "series": "Payment System Indicators monthly release workbook",
            "source_period": month, "source_url": RELEASE_ARCHIVE_URL,
            "downloaded_file": f"source_archive/rbi/xlsx/{path.name}",
            "retrieved_utc": today, "sha256": digest,
            "status": "browser_download_archived",
            "notes": "Own-month column extracted by scripts/import_rbi_psi_workbooks.py.",
        }
    with MANIFEST.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(current.values())


def main() -> None:
    workbooks = workbook_months()
    if not workbooks:
        raise SystemExit(f"No RBI PSI workbooks found in {WORKBOOKS}. Save the monthly "
                         "XLSX files there from a browser (see the script docstring).")
    cards = read_year_files("rbi_psi_card_ppi_own_month.csv")
    other = read_year_files("rbi_psi_other_rails.csv")
    infra = read_year_files("rbi_psi_infrastructure.csv")
    card_fields = ["month"] + [f"{f}_{u}" for f in CARD_PPI_FIELDS
                               for u in ("volume_lakh", "value_crore")] + [
        "rbi_release_id", "rbi_release_page_url", "official_xlsx_url"]
    other_fields = ["month"] + [f"{f}_{u}" for f in OTHER_RAIL_FIELDS
                                for u in ("volume_lakh", "value_crore")] + ["source_file"]
    infra_fields = ["month"] + [f for f, _ in INFRA_ITEMS] + ["unit", "source_file"]
    checks: list[dict[str, object]] = []
    for month, (path, _) in sorted(workbooks.items()):
        found, missing = extract(path, month)
        if missing:
            print(f"{month} {path.name}: missing {', '.join(missing)}", file=sys.stderr)
        existing = cards.get(month)
        if existing:
            for field, value in found.items():
                if field in existing and existing[field] != "":
                    old = float(existing[field])
                    tolerance = 1.0 if field.endswith("crore") else 0.011
                    status = "match" if abs(old - value) <= tolerance else "MISMATCH"
                    checks.append({"month": month, "field": field, "transcribed": old,
                                   "workbook": value, "status": status})
        elif all(f"{f}_volume_lakh" in found for f in CARD_PPI_FIELDS):
            cards[month] = {"month": month, **found,
                            "rbi_release_page_url": RELEASE_ARCHIVE_URL,
                            "official_xlsx_url": f"source_archive/rbi/xlsx/{path.name}"}
        if all(f"{f}_volume_lakh" in found for f in OTHER_RAIL_FIELDS):
            other[month] = {"month": month, **found, "source_file": path.name}
        if all(f in found for f, _ in INFRA_ITEMS):
            infra[month] = {"month": month, **{f: found[f] for f, _ in INFRA_ITEMS},
                            "unit": "lakh", "source_file": path.name}
    write_year_files("rbi_psi_card_ppi_own_month.csv", cards, card_fields)
    if other:
        write_year_files("rbi_psi_other_rails.csv", other, other_fields)
    if infra:
        write_year_files("rbi_psi_infrastructure.csv", infra, infra_fields)
    with CROSSCHECK.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=["month", "field", "transcribed",
                                                    "workbook", "status"])
        writer.writeheader()
        writer.writerows(checks)
    update_manifest(workbooks)
    mismatches = sum(1 for item in checks if item["status"] == "MISMATCH")
    print(f"Read {len(workbooks)} workbook(s); card/PPI months {len(cards)}, other-rail "
          f"months {len(other)}, infrastructure months {len(infra)}; "
          f"{len(checks)} cross-checks, {mismatches} mismatches (see {CROSSCHECK.name}).")
    if mismatches:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
